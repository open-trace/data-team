"""Build OpenTrace_Intermediate_Entity_Dictionary.xlsx (branded workbook).

Usage:
  python data-eng/scripts/build_opentrace_intermediate_entity_dictionary.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "OpenTrace_Intermediate_Entity_Dictionary.xlsx"
OUT_YAML = ROOT / "docs" / "intermediate_entity_dictionary_seed.yaml"
DOCS = ROOT / "docs"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from intermediate_dictionary_columns import (  # noqa: E402
    build_entities,
    collect_column_rows,
    parse_relationships,
    scan_sql_models,
)
from intermediate_dictionary_data import (  # noqa: E402
    COMMON_COLUMN_DESC,
    seed_as_dict,
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, max_width: int = 56) -> None:
    for idx, col in enumerate(ws.columns, start=1):
        length = 0
        for cell in col[:120]:
            length = max(length, len(str(cell.value or "")))
        ws.column_dimensions[get_column_letter(idx)].width = min(max_width, max(12, length + 2))


def _write(ws, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header(ws)
    _autosize(ws)
    # Wrap description-like columns
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column >= 4:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def _enable_filter(ws) -> None:
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions


def build() -> Path:
    DOCS.mkdir(parents=True, exist_ok=True)
    sql_models = scan_sql_models()
    entities = build_entities(sql_models)
    col_headers, col_rows, _ = collect_column_rows(entities=entities, sql_models=sql_models)

    blank = [r for r in col_rows if not str(r[4] or "").strip()]
    if blank:
        raise SystemExit(f"{len(blank)} columns missing descriptions (first={blank[0][:2]})")

    seed = seed_as_dict(entities)
    OUT_YAML.write_text(
        yaml.safe_dump(seed, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )

    wb = Workbook()

    # 00_README
    ws0 = wb.active
    ws0.title = "00_README"
    for line in [
        ["OpenTrace intermediate_dev — entity dictionary"],
        [],
        [
            "Complete catalogue of intermediate_dev dbt models and columns.\n"
            "Physical dataset: intermediate_dev (BQ_DATASET_INTERMEDIATE).\n"
            "Layer role: staging_dev → intermediate_dev (conformed spines + facts) → mart_dev.\n"
            "Every column has a detailed description derived from schema.yml, shared maps, and model context."
        ],
        [],
        ["Version", "Sep 2026"],
        ["Live models", str(len(entities))],
        ["Column rows", str(len(col_rows))],
        ["Regenerate", "python scripts/build_opentrace_intermediate_entity_dictionary.py"],
        [],
        ["Related"],
        ["OpenTrace_Mart_Entity_Dictionary.xlsx", "mart_dev branded dictionary"],
        ["ARCHITECTURE.md / GLOBAL_ARCHITECTURE.md", "Layer topology"],
        [],
        ["Sheets"],
        ["01_Table_Catalogue", "All int_* tables: domain, grain, purpose, upstream, joins"],
        ["02_Common_Columns", "Shared intermediate / ACF-style column cheat-sheet"],
        ["03_Domain_Map", "Domain folder → tables"],
        ["04_Columns", "Full per-table columns + detailed descriptions"],
        ["05_Relationships", "schema.yml + heuristic geo/source joins"],
        [
            "06_Data_Entity_Dictionary",
            "Staging-style flat sheet: Intermediate Model | Upstream Table | Entity/Column | Description",
        ],
    ]:
        ws0.append(line)
    ws0["A1"].font = Font(bold=True, size=14)
    _autosize(ws0, max_width=72)

    # 01_Table_Catalogue
    ws1 = wb.create_sheet("01_Table_Catalogue")
    cat_headers = [
        "table_name",
        "domain",
        "grain",
        "purpose",
        "primary_key",
        "upstream",
        "joins",
        "path",
    ]
    cat_rows: list[list[object]] = []
    for e in entities:
        cat_rows.append(
            [
                e["table_name"],
                e["domain"],
                e["grain"],
                e["purpose"],
                e.get("primary_key") or "",
                e.get("upstream") or "",
                e.get("joins") or "",
                e.get("path") or "",
            ]
        )
    _write(ws1, cat_headers, cat_rows)
    _enable_filter(ws1)

    # 02_Common_Columns
    ws2 = wb.create_sheet("02_Common_Columns")
    common_rows = [
        [k, v.get("role") or "", v.get("description") or "", v.get("example") or "", v.get("data_type") or ""]
        for k, v in sorted(COMMON_COLUMN_DESC.items())
    ]
    _write(
        ws2,
        ["column_name", "role", "description", "example", "data_type"],
        common_rows,
    )
    _enable_filter(ws2)

    # 03_Domain_Map
    ws3 = wb.create_sheet("03_Domain_Map")
    by_domain: dict[str, list[str]] = {}
    for e in entities:
        by_domain.setdefault(str(e["domain"]), []).append(str(e["table_name"]))
    domain_rows: list[list[object]] = []
    for domain in sorted(by_domain):
        tables = sorted(by_domain[domain])
        domain_rows.append(
            [
                domain,
                len(tables),
                ", ".join(tables),
                f"Intermediate models under dbt/models/intermediate_dev/{domain}/",
            ]
        )
    _write(ws3, ["domain", "table_count", "tables", "notes"], domain_rows)
    _enable_filter(ws3)

    # 04_Columns
    ws4 = wb.create_sheet("04_Columns")
    _write(ws4, col_headers, col_rows)
    _enable_filter(ws4)

    # 05_Relationships
    ws5 = wb.create_sheet("05_Relationships")
    # Only emit geo/source heuristics when the column actually exists
    cols_by_table: dict[str, set[str]] = {
        name: set(info.get("columns") or []) for name, info in sql_models.items()
    }
    rel_rows: list[list[str]] = []
    for r in parse_relationships(entities):
        table, col = r[0], r[1]
        if col in {"geo_key", "source_natural_key"} and col not in cols_by_table.get(table, set()):
            continue
        rel_rows.append(r)
    _write(
        ws5,
        ["from_table", "from_column", "to_table", "to_column", "notes"],
        rel_rows,
    )
    _enable_filter(ws5)

    # 06_Data_Entity_Dictionary — staging-style four columns
    ws6 = wb.create_sheet("06_Data_Entity_Dictionary")
    flat_headers = [
        "Intermediate Model",
        "Upstream Table",
        "Entity / Column",
        "Description",
    ]
    entities_by_name = {e["table_name"]: e for e in entities}
    flat_rows: list[list[object]] = []
    for table_name, col_name, _dtype, _role, desc, _ex in col_rows:
        ent = entities_by_name.get(str(table_name), {})
        flat_rows.append(
            [
                table_name,
                ent.get("upstream") or "",
                col_name,
                desc,
            ]
        )
    _write(ws6, flat_headers, flat_rows)
    _enable_filter(ws6)
    # Emphasize this sheet as the staging analogue
    for row in ws6.iter_rows(min_row=2, max_row=ws6.max_row, min_col=4, max_col=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws6.column_dimensions["D"].width = 72

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} ({len(entities)} models, {len(col_rows)} columns)")
    print(f"Wrote {OUT_YAML}")
    return OUT


if __name__ == "__main__":
    build()
