"""Build the landing datapoint report with real country names, charts, and a PDF."""

import csv
import re
import subprocess
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import yaml
from docx import Document
from docx.shared import Inches, Pt

SCRIPT_DIR = Path(__file__).resolve().parent
OUT_DIR = SCRIPT_DIR.parent / "out" / "dashboard"
DETAIL_PATH = OUT_DIR / "datapoints_by_class_country_year.csv"
CATALOG_PATH = OUT_DIR / "indicator_class_catalog.csv"
NAMES_PATH = SCRIPT_DIR / "country_names.yaml"
RESOLVED_PATH = OUT_DIR / "datapoints_by_class_country_resolved.csv"
DOCX_PATH = OUT_DIR / "landing_datapoint_report.docx"
PDF_PATH = OUT_DIR / "landing_datapoint_report.pdf"

EXPECTED_TOTAL = 137_740_480
UNRESOLVED = "No resolved country"
UNKNOWN_YEAR = "unknown"
CHART_COUNTRY_LIMIT = 20
_NUMBER = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
_APOSTROPHES = str.maketrans({"’": "'", "‘": "'", "`": "'", "´": "'", "ʼ": "'"})


def repair_text(value: str) -> str:
    for encoding in ("latin1", "cp1252"):
        try:
            fixed = value.encode(encoding).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
        if fixed:
            return fixed
    return value


def norm_key(value: str) -> str:
    text = repair_text(value).replace("_", " ").translate(_APOSTROPHES)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"\s+", " ", text.casefold().strip())
    return text


def is_number(value: str) -> bool:
    return bool(_NUMBER.fullmatch(value.strip().replace(",", "")))


def load_country_lookup(path: Path) -> dict[str, str]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    lookup: dict[str, str] = {}

    def add(label: str, country: str) -> None:
        key = norm_key(label)
        if not key:
            return
        previous = lookup.get(key)
        if previous is not None and previous != country:
            raise ValueError(f"{label!r} matches both {previous} and {country}")
        lookup[key] = country

    for row in data.get("countries") or []:
        name = str(row["name"]).strip()
        add(name, name)
        add(str(row["iso2"]), name)
        add(str(row["iso3"]), name)
        for alias in row.get("aliases") or []:
            add(str(alias), name)
    return lookup


def resolve_country(value: str, lookup: dict[str, str]) -> str:
    raw = (value or "").strip()
    if not raw or is_number(raw):
        return UNRESOLVED
    key = norm_key(raw)
    if not key or key in {"unknown", "none", "null", "nan"}:
        return UNRESOLVED
    return lookup.get(key, UNRESOLVED)


def clean_year(value: str) -> str:
    text = (value or "").strip()
    if not text or text.casefold() == UNKNOWN_YEAR:
        return UNKNOWN_YEAR
    if re.fullmatch(r"\d{4}\.0", text):
        return text[:-2]
    return text


def year_sort_key(year: str) -> tuple:
    if year.isdigit():
        return (0, int(year), "")
    match = re.match(r"(\d{4})", year)
    if match:
        return (1, int(match.group(1)), year)
    return (2, 0, year)


def load_classes(path: Path) -> dict[str, dict[str, str]]:
    classes: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            class_id = row["class_id"]
            if class_id not in classes:
                classes[class_id] = {
                    "name": row["class_name"],
                    "definition": row["opentrace_definition"],
                }
    return classes


def load_counts(path: Path, lookup: dict[str, str]) -> dict[str, dict[str, Counter[str]]]:
    counts: dict[str, dict[str, Counter[str]]] = defaultdict(lambda: defaultdict(Counter))
    with path.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            country = resolve_country(row["country"], lookup)
            year = clean_year(row["year"])
            counts[row["class_id"]][country][year] += int(row["n"])
    return counts


def class_total(countries: dict[str, Counter[str]]) -> int:
    return sum(sum(years.values()) for years in countries.values())


def known_years(countries: dict[str, Counter[str]]) -> Counter[str]:
    totals: Counter[str] = Counter()
    for years in countries.values():
        for year, count in years.items():
            if year != UNKNOWN_YEAR:
                totals[year] += count
    return totals


def resolved_only(countries: dict[str, Counter[str]]) -> dict[str, Counter[str]]:
    return {country: years for country, years in countries.items() if country != UNRESOLVED}


def write_resolved_csv(
    path: Path,
    counts: dict[str, dict[str, Counter[str]]],
    class_order: list[str],
) -> None:
    rows = []
    for class_id in class_order:
        countries = counts[class_id]
        names = sorted(countries, key=lambda name: (name == UNRESOLVED, name.casefold()))
        for country in names:
            years = countries[country]
            for year in sorted(years, key=lambda item: (item == UNKNOWN_YEAR, year_sort_key(item))):
                rows.append({
                    "class_id": class_id,
                    "country": country,
                    "year": year,
                    "n": years[year],
                })
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["class_id", "country", "year", "n"])
        writer.writeheader()
        writer.writerows(rows)


def _axis_ticks(value: float, _pos: int) -> str:
    number = float(value)
    if abs(number) >= 1_000_000:
        return f"{number / 1_000_000:.1f}M"
    if abs(number) >= 1_000:
        return f"{number / 1_000:.0f}k"
    return f"{number:.0f}"


def save_barh(path: Path, labels: list[str], values: list[int], title: str, xlabel: str) -> None:
    fig, ax = plt.subplots(figsize=(8, max(2.8, 0.38 * len(labels) + 1.2)))
    ax.barh(labels[::-1], values[::-1], color="#1f4e79")
    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.xaxis.set_major_formatter(mticker.FuncFormatter(_axis_ticks))
    fig.tight_layout()
    fig.savefig(path, dpi=130, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_year_bars(path: Path, labels: list[str], values: list[int], title: str) -> None:
    width = min(14, max(8, 0.28 * len(labels) + 2))
    fig, ax = plt.subplots(figsize=(width, 4.2))
    ax.bar(labels, values, color="#3d6b8c")
    ax.set_title(title)
    ax.set_ylabel("Datapoints")
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(_axis_ticks))
    ax.tick_params(axis="x", labelrotation=90, labelsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=130, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def _comma(value: int) -> str:
    return f"{value:,}"


def _add_table(doc: Document, headers: list[str], records: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for record in records:
        cells = table.add_row().cells
        for index, value in enumerate(record):
            cells[index].text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def _class_label(class_id: str, classes: dict[str, dict[str, str]]) -> str:
    name = (classes.get(class_id) or {}).get("name") or "Unclassified"
    return f"{class_id} — {name}"


def build_document(
    path: Path,
    counts: dict[str, dict[str, Counter[str]]],
    classes: dict[str, dict[str, str]],
    class_order: list[str],
    grand: int,
    *,
    title: str = "Landing datapoints by indicator class, country, and year",
    source_label: str = "landing tables",
    expected_total: int | None = EXPECTED_TOTAL,
) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    section = doc.sections[0]
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)

    resolved_total = 0
    country_names: set[str] = set()
    for countries in counts.values():
        for country, years in countries.items():
            amount = sum(years.values())
            if country != UNRESOLVED:
                resolved_total += amount
                country_names.add(country)

    doc.add_heading(title, level=0)
    doc.add_paragraph(
        f"The {source_label} hold {_comma(grand)} datapoints. "
        f"{_comma(resolved_total)} of them sit in {len(country_names)} resolved countries. "
        f"{_comma(grand - resolved_total)} stay under {UNRESOLVED}: "
        "the source value was a coordinate, a city or trial site, a region, or blank."
    )
    doc.add_paragraph(
        "ISO codes and spelling variants are shown as one English country name. "
        "NGA, NG, and nigeria are Nigeria. Cities and trial sites, such as Ubiaja, stay out of the country list."
    )

    class_labels = [_class_label(class_id, classes) for class_id in class_order]
    class_values = [class_total(counts[class_id]) for class_id in class_order]
    with TemporaryDirectory() as tmp_name:
        tmp = Path(tmp_name)
        overview = tmp / "classes.png"
        save_barh(overview, class_labels, class_values, "Datapoints by indicator class", "Datapoints")
        doc.add_picture(str(overview), width=Inches(6.5))

        overview_rows = []
        for class_id in class_order:
            countries = counts[class_id]
            resolved = resolved_only(countries)
            years = known_years(countries)
            overview_rows.append([
                _class_label(class_id, classes),
                _comma(class_total(countries)),
                str(len(resolved)),
                str(len(years)),
            ])
        _add_table(
            doc,
            ["Indicator class", "Datapoints", "Countries", "Years"],
            overview_rows,
        )

        for class_id in class_order:
            countries = counts[class_id]
            resolved = resolved_only(countries)
            total = class_total(countries)
            unresolved_n = sum(resolved_years.values()) if (resolved_years := countries.get(UNRESOLVED)) else 0
            years = known_years(countries)
            unknown_n = sum(bucket[UNKNOWN_YEAR] for bucket in countries.values())
            ordered_countries = sorted(resolved, key=lambda name: (-sum(resolved[name].values()), name.casefold()))
            ordered_years = sorted(years, key=year_sort_key)
            label = _class_label(class_id, classes)
            definition = (classes.get(class_id) or {}).get("definition") or ""

            doc.add_page_break()
            doc.add_heading(label, level=1)
            if definition:
                doc.add_paragraph(definition)
            doc.add_paragraph(f"Datapoints: {_comma(total)}.")
            if ordered_countries:
                doc.add_paragraph(
                    f"Resolved countries: {len(ordered_countries)}. "
                    + ", ".join(ordered_countries)
                    + "."
                )
            else:
                doc.add_paragraph("Resolved countries: 0.")
            if ordered_years:
                doc.add_paragraph(
                    f"Years: {len(ordered_years)}. " + ", ".join(ordered_years) + "."
                )
            else:
                doc.add_paragraph("Years: none recorded.")
            if unknown_n:
                doc.add_paragraph(f"Year not recorded: {_comma(unknown_n)} datapoints.")
            if unresolved_n:
                doc.add_paragraph(f"{UNRESOLVED}: {_comma(unresolved_n)} datapoints.")

            if ordered_countries:
                chart_names = ordered_countries[:CHART_COUNTRY_LIMIT]
                chart_values = [sum(resolved[name].values()) for name in chart_names]
                country_chart = tmp / f"{class_id}_countries.png"
                title = f"{class_id} datapoints by country"
                if len(ordered_countries) > CHART_COUNTRY_LIMIT:
                    title += f" (largest {CHART_COUNTRY_LIMIT})"
                save_barh(country_chart, chart_names, chart_values, title, "Datapoints")
                doc.add_picture(str(country_chart), width=Inches(6.5))
                if len(ordered_countries) > CHART_COUNTRY_LIMIT:
                    doc.add_paragraph(
                        f"The chart shows the {CHART_COUNTRY_LIMIT} countries with the most datapoints. "
                        "The sections below list every resolved country."
                    )
            if ordered_years:
                year_chart = tmp / f"{class_id}_years.png"
                save_year_bars(
                    year_chart,
                    ordered_years,
                    [years[year] for year in ordered_years],
                    f"{class_id} datapoints by year",
                )
                doc.add_picture(str(year_chart), width=Inches(6.5))

            for country in ordered_countries:
                year_counts = resolved[country]
                known = {year: count for year, count in year_counts.items() if year != UNKNOWN_YEAR}
                missing = year_counts.get(UNKNOWN_YEAR, 0)
                country_years = sorted(known, key=year_sort_key)
                doc.add_heading(country, level=2)
                doc.add_paragraph(
                    f"{_comma(sum(year_counts.values()))} datapoints in {label}. "
                    f"Years covered: {len(country_years)}."
                )
                if country_years:
                    _add_table(
                        doc,
                        ["Year", "Datapoints"],
                        [[year, _comma(known[year])] for year in country_years],
                    )
                if missing:
                    doc.add_paragraph(f"Year not recorded: {_comma(missing)} datapoints.")

    doc.add_page_break()
    doc.add_heading("Total check", level=1)
    doc.add_paragraph(
        f"The indicator classes, including datapoints with no resolved country, "
        f"sum to {_comma(grand)}."
    )
    if expected_total is not None and grand != expected_total:
        doc.add_paragraph(
            f"The landing total from the source count is {_comma(expected_total)}."
        )
    doc.save(path)


def export_pdf(docx: Path, pdf: Path) -> None:
    docx_s = str(docx.resolve()).replace("'", "''")
    pdf_s = str(pdf.resolve()).replace("'", "''")
    command = (
        "$ErrorActionPreference = 'Stop'; "
        "$word = New-Object -ComObject Word.Application; "
        "try { "
        "$word.Visible = $false; "
        "$word.DisplayAlerts = 0; "
        f"$doc = $word.Documents.Open('{docx_s}'); "
        f"$doc.SaveAs2('{pdf_s}', 17); "
        "$doc.Close($false); "
        "} finally { $word.Quit(); }"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", command], check=True)


def main() -> int:
    if not DETAIL_PATH.exists():
        print(f"Missing {DETAIL_PATH}", file=sys.stderr)
        return 1
    lookup = load_country_lookup(NAMES_PATH)
    classes = load_classes(CATALOG_PATH)
    counts = load_counts(DETAIL_PATH, lookup)
    grand = sum(class_total(countries) for countries in counts.values())
    class_order = sorted(counts, key=lambda class_id: (-class_total(counts[class_id]), class_id))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_resolved_csv(RESOLVED_PATH, counts, class_order)
    print(f"Grand total: {grand}")
    if grand != EXPECTED_TOTAL:
        print(f"Expected {EXPECTED_TOTAL}", file=sys.stderr)
        return 1
    build_document(DOCX_PATH, counts, classes, class_order, grand)
    print(f"Wrote {DOCX_PATH}")
    export_pdf(DOCX_PATH, PDF_PATH)
    print(f"Wrote {PDF_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
