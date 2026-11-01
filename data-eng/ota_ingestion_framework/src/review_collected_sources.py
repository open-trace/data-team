# Author: Daniel Wondyfiraw
# Updated: 2026-10-30
"""Review a completed handover offline; never fetch, delete or change collection rules.

Run: python -m src.review_collected_sources
Recommendations describe saved evidence. Unknown sources remain unreviewed rather
than being declared irrelevant from their URL. Generic tables still need source
and business validation before reporting use.
"""
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.common.source_config import DATA, GENERATED
from src.common.logging import get_logger

logger=get_logger(__name__)
PRIORITY_PREFIXES=('za_', 'ke_', 'ng_', 'global_world_bank_')


def inspect_export(path, item):
    with path.open(encoding='utf-8-sig',newline='') as stream:
        reader=csv.DictReader(stream)
        columns=reader.fieldnames or []
        rows=list(reader)
    sid=item['source_id']
    sample=rows[:20]
    # Generic grids have a publisher header row inside their first data row.
    first=list(rows[0].values())[1:] if rows and columns[0] in {'row_number','excel_row'} else columns
    header=' | '.join(str(x) for x in first).lower()
    preview=' '.join(str(v) for row in sample for v in row.values()).lower()
    category='REVIEW_REQUIRED';reason='Saved table requires inspection of its meaning, period and units.'
    dates=[];price_rows=0
    if sid.startswith('api_api_ninjas_com_v1_commodityprice_name_') and {'price','unit','currency_unit','updated'}.issubset(columns):
        valid=[r for r in rows if r.get('price') and r.get('unit') and r.get('currency_unit') and r.get('updated')]
        if valid:
            category='BENCHMARK_QUOTES';reason='Structured financial quote fields present; retain native currency/unit and contract basis.'
            price_rows=len(valid)
            for row in valid:
                try: dates.append(datetime.fromtimestamp(int(row['updated']),timezone.utc).date().isoformat())
                except (ValueError,OverflowError,OSError): pass
    elif {'county','commodity','priceDate'}.issubset(columns):
        category='COUNTY_PRICE_SUMMARIES';reason='County-level commodity summaries; retain missing-price entries and source periods; currency and aggregation need review.'
        price_rows=sum(any(r.get(k) for k in ['retailPrice','wholesalePrice','displayPrice']) for r in rows)
        for row in rows:
            try:dates.append(datetime.strptime(row.get('priceDate',''),'%Y-%m-%d').date().isoformat())
            except ValueError:pass
    elif {'desc','averageprice','source_observation_date'}.issubset(columns):
        category='MARKET_PRICE_ROWS';reason='Commodity/package price rows with publisher observation dates; verify currency and basis.'
        price_rows=sum(bool(r.get('desc') and r.get('averageprice')) for r in rows)
        for row in rows:
            try:dates.append(datetime.strptime(row.get('source_observation_date',''),'%Y-%m-%d').date().isoformat())
            except ValueError:pass
    elif urlsplit(item.get('source_url','')).hostname=='microdata.nigerianstat.gov.ng' and 'item label' in header and ('average of' in header or 'north central' in header):
        category='FOOD_PRICE_GRID';reason='NBS national or zonal food-price grid; source period is in the workbook header/label. Complete four-state coverage is not established.'
    elif {'commoditiesclicktosort','averageprice','date'}.issubset(columns):
        category='MARKET_PRICE_ROWS';reason='Commodity, low/average/high prices, package details and source date present; units/basis require mapping.'
        valid=[r for r in rows if r.get('commoditiesclicktosort') and r.get('date') and r.get('averageprice')]
        price_rows=len(valid)
        for row in valid:
            try:dates.append(datetime.strptime(row['date'],'%d/%b/%Y').date().isoformat())
            except ValueError: pass
    elif {'observation_date','commodity_source','average_price','market'}.issubset(columns):
        category='MARKET_PRICE_ROWS';reason='Structured market observations with source dates and package fields; retain quality findings and verify currency/price basis.'
        valid=[r for r in rows if r.get('commodity_source') and r.get('observation_date') and r.get('average_price')]
        price_rows=len(valid)
        for row in valid:
            try:dates.append(datetime.strptime(row['observation_date'],'%Y-%m-%d').date().isoformat())
            except ValueError:pass
    elif all(x in header for x in ['commodity','price','market','date']) and 'category' in header:
        category='MARKET_PRICE_TABLE';reason='Publisher table has commodity, category, price, market and date; validate publisher provenance and recency.'
        price_rows=max(0,len(rows)-1)
        for row in rows[1:]:
            value=list(row.values())[-1]
            try:dates.append(datetime.strptime(value,'%d %b %Y').date().isoformat())
            except ValueError:pass
    elif ('company' in header and 'mkt cap' in header) or ('stock' in header and 'share o/s' in header):
        category='NON_COMMODITY_TABLE';reason='Equity/company market-cap table; exclude from agricultural-price modelling.'
    elif ('ecole' in header or 'ecole' in preview or 'école' in header) and 'localisation' in header:
        category='NON_COMMODITY_TABLE';reason='School/training directory, not a commodity dataset.'
    elif 'express international import' in preview:
        category='NON_COMMODITY_TABLE';reason='Postal express object counts, not agricultural import quantities.'
    elif ('assistant director' in preview or 'principal scientific officer' in preview) and '@' in preview:
        category='NON_COMMODITY_TABLE';reason='Staff/contact directory; no commodity records verified.'
    elif 'auction date' in header and 'venue' in header:
        category='CATALOGUE_OR_SCHEDULE';reason='Auction schedule; price results were not established by this table.'
    elif 'market name' in header and 'market days' in header:
        category='MARKET_DIRECTORY';reason='Market names, locations and trading days; useful geography reference, not price observations.'
    elif sid=='api_ninjas_com_api_commodityprice':
        category='API_DOCUMENTATION';reason='Documentation/examples and performance tables; use authenticated API URLs for actual quotes.'
    elif 'code' in header and 'contract' in header and 'available from' in header:
        category='COMMODITY_CATALOGUE';reason='Supported commodity/contract catalogue; not returned price observations.'
    elif 'start | end' in header or header in {'name | role','name | affiliation | role','name | grant number | role','name | abbreviation | grant number | role'}:
        category='CATALOGUE_METADATA';reason='Coverage/attribution catalogue metadata; no actual commodity records established.'
    elif sid=='www_namc_co_za' or sid=='www_namc_co_za_2':
        category='NON_COMMODITY_TABLE';reason='Saved table contains a meeting/event listing, not agricultural observations.'
    elif item['data_kind']=='document_text':
        category='DOCUMENT_REVIEW';reason='PDF text lines only; inspect document subject and period before extracting data.'
        if 'how to use rrs' in preview or 'atelier sur' in preview or 'data and software policy' in preview:
            category='NON_COMMODITY_DOCUMENT';reason='Saved text describes instructions, a workshop or software policy; no commodity records established.'
    elif sid=='www_mansamarkets_com_commodities_sesame_nigeria' and 'price' in header:
        category='PRICE_TABLE_CANDIDATE';reason='Published commodity/spot-price table; inspect source geography, price basis and observation date; related stock tables excluded separately.'
    elif 'export value' in header and ('volume' in header or 'world share' in header):
        category='TRADE_STATISTICS_CANDIDATE';reason='Export value/volume table; trade quantities and unit values differ from current wholesale prices.'
    elif 'market level' in header and 'price' in header:
        category='TRADE_UNIT_VALUE_CANDIDATE';reason='Export unit-value table, not verified local physical market prices.'
    elif item['data_kind']=='workbook_grid':
        category='WORKBOOK_REVIEW';reason='Raw workbook grid; inspect title, headers, units and periods before treating rows as observations.'
        if 'commodity price forecasts' in preview:
            category='FORECAST_GRID';reason='Commodity forecasts; keep separate from realised prices.'
        elif 'food price index' in preview or 'food price indices' in preview:
            category='PRICE_INDEX_GRID';reason='Food-price index grid; indices are not absolute local commodity prices.'
    return dict(category=category,reason=reason,rows=len(rows),price_rows=price_rows,
                latest_observation=max(dates,default=''),header=' | '.join(str(x) for x in first)[:250])


def main():
    base=DATA/'exports/latest'
    manifest=json.loads((base/'export_report.json').read_text())
    batch=json.loads((GENERATED/'batch_collection.json').read_text())
    inventory=list(csv.DictReader(sorted(GENERATED.glob('urls_*.csv'))[-1].open(encoding='utf-8-sig')))
    by_source=defaultdict(list);files=[];seen={};integrity=[]
    for item in manifest['files']:
        path=base/item['file']
        reviewed=inspect_export(path,item)
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        if digest!=item['sha256'] or reviewed['rows']!=item['rows']:
            integrity.append(item['file'])
        duplicate=seen.get(digest,'');seen.setdefault(digest,item['file'])
        row=dict(source_id=item['source_id'],file=item['file'],data_kind=item['data_kind'],**reviewed,
                 exact_same_file_as=duplicate)
        files.append(row);by_source[item['source_id']].append(row)
    status={s['source_id']:s for s in batch['sources']}
    sources=[]
    useful={'MARKET_PRICE_ROWS','MARKET_PRICE_TABLE','BENCHMARK_QUOTES','COUNTY_PRICE_SUMMARIES'}
    for entry in inventory:
        sid=entry['source_id'];exports=by_source[sid];counts=Counter(f['category'] for f in exports)
        priority=sid.startswith(PRIORITY_PREFIXES)
        if any(c in useful for c in counts):
            action='PRIORITISE_DATA_REVIEW';reason='Saved structured price/benchmark records; review dates, units, provenance and scope before reporting.'
        elif any(c in counts for c in ['PRICE_TABLE_CANDIDATE','PRICE_INDEX_GRID','FORECAST_GRID','TRADE_STATISTICS_CANDIDATE','TRADE_UNIT_VALUE_CANDIDATE','FOOD_PRICE_GRID']):
            action='REVIEW_RELATED_MEASURES';reason='Relevant candidate measures; separate trade, index, forecast and benchmark meanings.'
        elif priority:
            action='PRIORITY_SOURCE_INVESTIGATION';reason='Partner/benchmark candidate; current access or extraction outcome does not establish complete usable data.'
        elif exports and all(f['category'].startswith('NON_COMMODITY') or f['category']=='API_DOCUMENTATION' for f in exports):
            action='DEPRIORITISE_CURRENT_EXPORTS';reason='Inspected saved output is unrelated or documentation; this does not prove the whole website is irrelevant.'
        elif exports:
            action='REVIEW_SAVED_CONTENT';reason='Inspect saved grids, documents or catalogue material; no verified price dataset inferred.'
        else:
            action='UNASSESSED_RELEVANCE';reason='No exported records to judge; preserve endpoint for source-owner review.'
        sources.append(dict(source_id=sid,source_name=entry['display_name'],url=entry['source_url'],
            batch_status=status.get(sid,{}).get('status','NOT_IN_BATCH'),recommendation=action,
            files=len(exports),raw_rows=sum(f['rows'] for f in exports),categories='; '.join(f'{k}: {v}' for k,v in counts.items()),
            latest_observation=max((f['latest_observation'] for f in exports),default=''),reason=reason,
            batch_issue=status.get(sid,{}).get('reason','')))
    counts=Counter(f['source_id'] for f in manifest['files'])
    differences=[dict(source_id=s['source_id'],batch_files=len(s['csv_files']),handover_files=counts[s['source_id']],batch_rows=s['rows_exported'])
                 for s in batch['sources'] if len(s['csv_files'])!=counts[s['source_id']]]
    out=GENERATED/'relevance_review';out.mkdir(exist_ok=True)
    for name,rows in [('sources',sources),('exports',files)]:
        with (out/f'{name}.csv').open('w',encoding='utf-8-sig',newline='') as stream:
            w=csv.DictWriter(stream,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    summary=dict(reviewed_at_utc=datetime.now(timezone.utc).isoformat(),candidate_urls=len(sources),export_sources=sum(bool(v) for v in by_source.values()),
                 files=len(files),raw_rows=sum(f['rows'] for f in files),recommendations=dict(Counter(s['recommendation'] for s in sources)),
                 export_categories=dict(Counter(f['category'] for f in files)),exact_repeated_files=sum(bool(f['exact_same_file_as']) for f in files),integrity_failures=integrity,file_count_differences=differences,
                 scope='Offline evidence-based triage, not production approval. No network requests, records deleted, or collection exclusions applied. Unknown endpoints remain unassessed. Price-table rows require business validation.')
    (out/'summary.json').write_text(json.dumps(summary,indent=2)+'\n')
    book=Workbook();book.remove(book.active)
    for name,rows in [('Sources',sources),('Export files',files),('Priorities',[s for s in sources if s['recommendation'] in {'PRIORITISE_DATA_REVIEW','PRIORITY_SOURCE_INVESTIGATION','REVIEW_RELATED_MEASURES'}])]:
        sheet=book.create_sheet(name);sheet.append(list(rows[0]))
        for r in rows:sheet.append([r[k] for k in rows[0]])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.fill=PatternFill('solid',fgColor='14564B');c.font=Font(color='FFFFFF',bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(65,max(18,len(str(col[0].value))+4))
        for row in sheet.iter_rows():
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    book.save(out/'Source_Relevance_Review.xlsx')
    logger.info('Reviewed %s candidates and %s saved CSVs. Open %s',len(sources),len(files),out/'Source_Relevance_Review.xlsx')
    print(json.dumps(summary,indent=2))
    return 1 if integrity else 0


if __name__=='__main__':raise SystemExit(main())
