# Author: Daniel Wondyfiraw
# Updated: 2026-10-30
"""Source-specific diagnostics. This module never lands or writes source rows."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import datetime, timedelta
from itertools import islice
from urllib.parse import urljoin, urlsplit

import openpyxl

from src.common.fetching_method import fetch_json_records, fetch_naerls_selections
from src.common.source_testing import (
    Page, ProbeError, check_records, display_url, form_data, json_pages,
    normalise, numeric, observation_date, secure_url, table_records,
)


def html_page(body):
    return Page(body.decode("utf-8-sig", "replace"))


def json_response(body):
    try:
        value = json.loads(body)
    except (ValueError, UnicodeError):
        raise ProbeError("Expected JSON; actual response is not valid JSON.") from None
    if isinstance(value, dict) and (value.get("success") is False or value.get("error")):
        raise ProbeError("JSON response reports an application error.")
    return value


def discover(evidence, name, page, url):
    links = []
    for link in page.links:
        href = link["href"]
        if any(x in href.lower() for x in ("/download/", ".csv", ".xlsx", ".zip", ".pdf", "/api/", "get-microdata")):
            target = urljoin(url, href)
            if urlsplit(target).scheme in {"http", "https"}:
                links.append(display_url(target))
    evidence.check(name + ".discovery", "PASS", "Actual HTML inspected for forms and data/download links. Discovery does not verify those resources.", forms=[{"method": f["method"], "action": display_url(urljoin(url, f["action"])), "fields": [v.get("name") for v in f["inputs"] + f["selects"] if v.get("name")]} for f in page.forms], discovered_data_links=list(dict.fromkeys(links))[:30], data_links_count=len(set(links)), script_count=len(page.scripts))


def checked_html(client, evidence, name, url, *, read_limit=None):
    body, final = client.get(name, url, read_limit=read_limit)
    page = html_page(body)
    discover(evidence, name, page, final)
    return page, final


def post_form(client, name, page, current, field, values, *, read_limit=None):
    form = page.form(field)
    target = secure_url(urljoin(current, form["action"]))
    if urlsplit(target).netloc != urlsplit(current).netloc:
        raise ProbeError("Unexpected cross-host form action.")
    if form["method"] != "POST":
        raise ProbeError("Expected a POST form; the source has changed.")
    data = form_data(form)
    data.update(values)
    body, final = client.request(name, "POST", target, data=data, read_limit=read_limit)
    return html_page(body), final


def price_order(evidence, name, records, low, high, average):
    bad, checked = 0, 0
    for row in records:
        values = [numeric(row.get(k)) for k in (low, high, average)]
        if all(v is not None for v in values):
            checked += 1
            bad += not (0 <= values[0] <= values[2] <= values[1])
    evidence.check(name + ".price_order", "FAIL" if bad else "PASS" if checked else "NOT_TESTED", "Checked low <= average <= high using original price values.", rows_checked=checked, inconsistent_rows=bad)


def durban_table(evidence, name, page, profile, *, complete=True):
    headers, rows = table_records(page, ["Weight (Kg)", "Low Price", "High Price", "Average Price", "Date"])
    commodity = normalise(headers[0])
    check_records(evidence, name, rows, date_field="date", price_fields=["lowprice", "highprice", "averageprice"], commodity_field=commodity, max_age_days=profile.get("freshness_days"), complete=complete)
    price_order(evidence, name, rows, "lowprice", "highprice", "averageprice")
    return rows, commodity


def probe_durban(client, evidence, source, profile):
    # Read complete HTML within Client's existing byte/time limits. A 64 KB
    # prefix can cut off both table rows and commodity options.
    page, url = checked_html(client, evidence, "daily_table", source["source_url"])
    records, commodity_key = durban_table(evidence, "daily_table", page, profile, complete=evidence.report["endpoints"][-1]["complete"])
    select = next((s for s in page.form("descItem")["selects"] if s["name"] == "descItem"), None)
    if select is None:
        raise ProbeError("The commodity selector is missing from the returned form.")
    options = [{"label": " ".join(o["text"]).strip(), "value": o["value"]}
               for o in select["options"]]
    evidence.check("commodity_options", "PASS" if options else "FAIL",
                   "Commodity options read from the actual HTML form; these are labels, not verified price records.",
                   options=options, option_count=len(options))
    crops = profile.get("commodities")
    if crops is None:
        # Exercise at most two real options also present in the returned table.
        # Do not invent aliases such as TOMATOES when the form says TOMATO.
        present = {normalise(row[commodity_key]) for row in records}
        crops = list(dict.fromkeys(o["label"] for o in options
                                  if o["label"] and normalise(o["label"]) in present))[:2]
    evidence.check("commodity_filter_scope", "PASS" if crops else "NOT_TESTED",
                   "Only the listed form options are exercised; other options remain untested.",
                   selected_options=crops)
    for crop in crops:
        name = "commodity_filter_" + normalise(crop)
        try:
            option = next((o for o in select["options"] if normalise(" ".join(o["text"])) == normalise(crop)), None)
            if option is None:
                raise ProbeError("Configured commodity is missing from the actual form options.")
            value = option["value"] if option["value"] is not None else " ".join(option["text"])
            result, _ = post_form(client, name, page, url, "descItem", {"descItem": value})
            rows, key = durban_table(evidence, name, result, profile, complete=evidence.report["endpoints"][-1]["complete"])
            matches = bool(rows) and all(normalise(row[key]) == normalise(crop) for row in rows)
            evidence.check(name + ".filter", "PASS" if matches else "FAIL", "Every returned commodity checked against the submitted option.")
        except ProbeError as exc:
            evidence.check(name + ".filter", "FAIL", str(exc))
    try:
        dates = [observation_date(row.get("date")) for row in records]
        dates = [d for d in dates if d]
        if not dates or not profile.get("date_input_format"):
            raise ProbeError("A source observation date and reviewed date-input format are required.")
        target = max(dates) - timedelta(days=1)
        result, _ = post_form(client, "previous_date", page, url, "datepickers", {"datepickers": target.strftime(profile["date_input_format"]), "descItem": ""})
        _, rows = table_records(result, ["Weight (Kg)", "Low Price", "High Price", "Average Price", "Date"])
        if rows:
            evidence.check("previous_date.filter", "PASS" if all(observation_date(r.get("date")) == target for r in rows) else "FAIL", "Returned row dates compared with the requested date.", requested_date=target.isoformat(), record_count=len(rows))
        else:
            evidence.check("previous_date.filter", "NOT_TESTED", "The requested day returned no rows; date-filter correctness is unproven.", requested_date=target.isoformat())
    except ProbeError as exc:
        evidence.check("previous_date.filter", "NOT_TESTED", str(exc))
    evidence.check("pagination", "NOT_TESTED", "Complete HTML is requested, but server/client pagination and authoritative source totals have not been verified.", required=False)
    evidence.check("historical_coverage", "NOT_TESTED", "One previous-day query does not establish historical coverage.", required=False)


def select_event(page, preferred_mass=None):
    choices = []
    for row in page.rows:
        for link in row.links:
            match = re.search(r"__doPostBack\('([^']+)','(Select\$\d+)'\)", link)
            if match:
                choices.append((row, match.groups()))
    if not choices:
        raise ProbeError("No item-selection postback found in the actual response.")
    if preferred_mass is not None:
        preferred = next((c for c in choices if len(c[0].cells) == 7 and c[0].cells[3] == "1M" and numeric(c[0].cells[5]) == preferred_mass), None)
        if preferred:
            return preferred[1]
    return choices[0][1]


def probe_tshwane(client, evidence, source, profile):
    for crop in profile.get("commodities", ["TOMATO"]):
        name = normalise(crop)
        try:
            page, url = checked_html(client, evidence, name + ".form", source["source_url"])
            page, url = post_form(client, name + ".search", page, url, "__VIEWSTATE", {"ctl00$ContentPlaceHolder1$ddlMonth": "Contains", "ctl00$ContentPlaceHolder1$txtItem": crop, "ctl00$ContentPlaceHolder1$BtnSubmit": "Submit Query"})
            for stage in ("product", "package"):
                target, argument = select_event(page, 6 if crop == "TOMATO" and stage == "package" else None)
                page, url = post_form(client, name + "." + stage, page, url, "__VIEWSTATE", {"__EVENTTARGET": target, "__EVENTARGUMENT": argument})
            record = {normalise(row.cells[0]): row.cells[1] for row in page.rows if len(row.cells) == 2}
            found = re.search(r"Data as at\s+(\d{1,2}\s+\w+\s+\d{4})", " ".join(page.text), re.I)
            if not found or not observation_date(found.group(1)):
                raise ProbeError("Final statistics have no identifiable source observation date.")
            if normalise(crop) not in normalise(record.get("product")):
                raise ProbeError("Selected product differs from the requested commodity.")
            check_records(evidence, name, [record], date_field="date", date_override=found.group(1), price_fields=["lowestprice", "highestprice", "averageprice"], commodity_field="product", max_age_days=profile.get("freshness_days"))
            price_order(evidence, name, [record], "lowestprice", "highestprice", "averageprice")
            evidence.check(name + ".optional_fields", "PASS", "Blank source fields remain missing; no zero or price-per-kg value was invented.", missing_fields=[key for key in ("count", "voided", "randperkg") if not record.get(key)])
        except ProbeError as exc:
            evidence.check(name + ".records", "FAIL", str(exc))
    evidence.check("pagination", "NOT_APPLICABLE", "Product/package selection postbacks were tested; this profile does not declare paginated price pages.")
    evidence.check("commodity_coverage", "NOT_TESTED", "Only configured sample commodities/items are checked, not every listed package.", required=False)
    evidence.check("historical_coverage", "NOT_TESTED", "Daily item flow does not establish historical retrieval.", required=False)


def probe_joburg(client, evidence, source, profile):
    page, url = checked_html(client, evidence, "daily_table", source["source_url"])
    table = None
    for candidate in page.tables:
        for index, header in enumerate(candidate):
            keys = [normalise(x) for x in header.cells]
            value = next((k for k in keys if k.startswith("totalvaluesold")), None)
            mass = next((k for k in keys if k.startswith("totalkgsold") or k.startswith("totalkilogramssold")), None)
            commodity = next((k for k in keys if k.startswith("commodity")), None)
            if value and mass and commodity:
                table = (keys, candidate[index + 1:], value, mass, commodity)
                break
        if table:
            break
    if not table:
        raise ProbeError("Expected Joburg commodity/sales/kilogram table is missing.")
    keys, data_rows, value_key, mass_key, commodity_key = table
    rows = [dict(zip(keys, row.cells)) for row in data_rows if len(row.cells) == len(keys)]
    text = " ".join(page.text)
    matched = re.search(r"(?:data\s+(?:for|as\s+at)|prices\s+(?:for|as\s+at)|this\s+information\s+is\s+for)\s*:?\s*(\d{1,2}\s+\w+\s+\d{4})", text, re.I)
    observed = observation_date(matched.group(1)) if matched else None
    if observed is None:
        raise ProbeError("Observation-date label was not found; the retrieval date will not be substituted.")
    valid_sales = sum(
        bool(row.get(commodity_key, "").strip())
        and all(v is not None and v >= 0 for v in (numeric(row.get(value_key)), numeric(row.get(mass_key))))
        for row in rows
    )
    evidence.check("daily_sales.sales", "PASS" if rows and valid_sales == len(rows) else "FAIL", "Actual sales values and quantities inspected. These measures are not prices.", record_count=len(rows), valid_rows=valid_sales, invalid_rows=len(rows) - valid_sales, fields=keys, earliest_observation=observed.isoformat(), latest_observation=observed.isoformat())
    evidence.check("daily_prices.records", "NOT_TESTED", "This response publishes sales totals, not low/high/average prices. No price records or derived value/kg prices were verified.")
    for crop in profile.get("commodities", []):
        try:
            body, _ = client.get("commodity_filter_" + normalise(crop), url, params={"commodity": crop})
            result = html_page(body)
            error = any(x in " ".join(result.text).lower() for x in ("uncaught", "fatal error", "typeerror"))
            evidence.check("commodity_filter_" + normalise(crop), "FAIL" if error else "NOT_TESTED", "Server-side application error in HTTP 200." if error else "Response received; filtered table semantics need a dedicated check.")
        except ProbeError as exc:
            evidence.check("commodity_filter_" + normalise(crop), "FAIL", str(exc))
    evidence.check("pagination", "NOT_APPLICABLE", "Configured request retrieves the daily HTML table.")


def probe_strathmore(client, evidence, source, profile):
    root = source["source_url"]
    try:
        _, root = checked_html(client, evidence, "dashboard", root)
    except ProbeError as exc:
        evidence.check("dashboard", "FAIL", str(exc))
    paths = profile["endpoints"]
    options = None
    try:
        options_body, _ = client.get("options", urljoin(root, paths["options"]))
        options = json_response(options_body)
        if not isinstance(options, dict) or not isinstance(options.get("commodities"), list) or not isinstance(options.get("counties"), list):
            options = None
            raise ProbeError("Options endpoint does not provide commodity and county lists.")
        evidence.check("options.schema", "PASS", "Commodity and county option lists verified.")
    except ProbeError as exc:
        evidence.check("options.schema", "FAIL", str(exc))
    for endpoint in ("national", "county"):
        try:
            body, _ = client.get(endpoint, urljoin(root, paths[endpoint]))
            data = json_response(body)
            records = data.get("latestPrices") if isinstance(data, dict) else None
            if not isinstance(records, list):
                raise ProbeError("Summary endpoint lacks its expected latestPrices list.")
            check_records(evidence, endpoint, records, date_field="lastUpdated", price_fields=profile[endpoint + "_price_fields"], commodity_field="commodity", location_field="market", unit_field="unit", max_age_days=profile.get("freshness_days"))
        except ProbeError as exc:
            evidence.check(endpoint + ".records", "FAIL", str(exc))
    counties = profile["counties"]
    for crop in profile["commodities"]:
        name = "comparison_" + normalise(crop)
        try:
            if options is not None and (crop not in options["commodities"] or any(c not in options["counties"] for c in counties)):
                raise ProbeError("Requested commodity/county is not advertised by the options endpoint.")
            body, _ = client.get(name, urljoin(root, paths["comparison"]), params={"commodity": crop, "countyA": counties[0], "countyB": counties[1]})
            data = json_response(body)
            if not isinstance(data, dict) or data.get("commodity") != crop:
                raise ProbeError("Comparison returned the wrong commodity or response shape.")
            records = []
            for key, county in zip(("countyA", "countyB"), counties):
                original = data.get(key)
                if not isinstance(original, dict) or original.get("county") != county:
                    raise ProbeError("Comparison returned a missing or mismatched county.")
                record = dict(original)
                record["commodity"] = data["commodity"]
                records.append(record)
            check_records(evidence, name, records, date_field="priceDate", price_fields=["wholesalePrice", "retailPrice"], commodity_field="commodity", location_field="county", unit_field="unit", max_age_days=profile.get("freshness_days"), required_locations={c: [c] for c in counties})
        except ProbeError as exc:
            evidence.check(name + ".records", "FAIL", str(exc))
    evidence.check("pagination", "NOT_APPLICABLE", "Configured JSON routes return summaries or a fixed county pair; no next-link contract is declared.")
    evidence.check("provenance_and_terms", "NOT_TESTED", "University-published quotes still require original-provider and reuse review.", required=False)


def probe_csv(client, evidence, source, profile):
    body, _ = client.get("csv_download", source["source_url"], read_limit=profile.get("sample_bytes"))
    if body.lstrip().startswith((b"<", b"{")):
        raise ProbeError("Expected CSV, but the response is HTML/JSON rather than CSV records.")
    complete = evidence.report["endpoints"][-1]["complete"]
    try:
        text = body.decode("utf-8-sig")
        if not complete:
            text = text.rsplit("\n", 1)[0] + "\n"
        reader = csv.DictReader(io.StringIO(text), strict=True)
        columns = reader.fieldnames or []
        required = profile["required_fields"]
        if not set(required).issubset(columns):
            raise ProbeError("CSV schema does not contain all required source fields.")
        rows = [row for row in reader if not str(row.get(profile["date_field"], "")).startswith("#")]
    except (UnicodeError, csv.Error):
        raise ProbeError("CSV could not be decoded or parsed.") from None
    check_records(evidence, "csv", rows, date_field=profile["date_field"], price_fields=profile["price_fields"], commodity_field=profile["commodity_field"], location_field=profile.get("location_field"), unit_field=profile.get("unit_field"), max_age_days=profile.get("freshness_days"), required_locations=profile.get("required_locations"), complete=complete)
    evidence.check("pagination", "NOT_APPLICABLE", "CSV file response; byte-sample completeness is reported separately from API pagination.")


def safe_archive(body, byte_limit):
    try:
        archive = zipfile.ZipFile(io.BytesIO(body))
        if sum(info.file_size for info in archive.infolist()) > byte_limit * 5:
            archive.close()
            raise ProbeError("Uncompressed archive exceeds the inspection limit.")
        return archive
    except zipfile.BadZipFile:
        raise ProbeError("Expected a ZIP/XLSX container; the response is not a valid archive.") from None


def inspect_workbook(evidence, name, body, profile):
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise ProbeError("Expected a genuine readable XLSX workbook.") from None
    try:
        evidence.check(name + ".format", "PASS", "XLSX opened in memory; no workbook or source rows saved.", sheets=workbook.sheetnames)
        if profile.get("workbook") == "nbs_food_watch":
            observations, periods, items = 0, set(), set()
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                if not header or normalise(header[0]) != "itemlabel":
                    continue
                columns = {}
                for index, value in enumerate(header):
                    match = re.fullmatch(r"Average of ([A-Za-z]+)-(\d{2})", str(value or "").strip())
                    if not match:
                        continue
                    for fmt in ("%b-%y", "%B-%y"):
                        try:
                            columns[index] = datetime.strptime("-".join(match.groups()), fmt).strftime("%Y-%m")
                            break
                        except ValueError:
                            continue
                for row in rows:
                    if not row or not str(row[0] or "").strip():
                        continue
                    for index, period in columns.items():
                        value = numeric(row[index]) if index < len(row) else None
                        if value is not None and value >= 0:
                            observations += 1
                            periods.add(period)
                            items.add(str(row[0]).strip())
            evidence.check(name + ".records", "PASS" if observations else "NOT_TESTED",
                           "National monthly averages read from explicit Average of Month-Year columns. State extrema and zonal averages are not state price panels.",
                           record_count=observations, commodity_count=len(items), geographic_scope="national_aggregate",
                           observation_precision="month", earliest_observation=min(periods) if periods else None,
                           latest_observation=max(periods) if periods else None)
            evidence.check(name + ".four_state_panel", "NOT_TESTED",
                           "This inspected workbook supplies national averages, state high/low summaries and zonal averages. It does not establish a complete FCT/Edo/Benue/Delta price panel.")
            return
        if profile.get("workbook") == "pink_sheet":
            if "Monthly Prices" not in workbook.sheetnames:
                raise ProbeError("Expected Monthly Prices sheet is missing.")
            periods, numeric_cells = [], 0
            for row in workbook["Monthly Prices"].iter_rows(values_only=True):
                match = re.fullmatch(r"(\d{4})M(0[1-9]|1[0-2])", str(row[0] or "").strip()) if row else None
                if match:
                    values = [numeric(v) for v in row[1:]]
                    count = sum(v is not None for v in values)
                    if count:
                        periods.append(f"{match[1]}-{match[2]}")
                        numeric_cells += count
            evidence.check(name + ".records", "PASS" if periods else "FAIL", "Monthly source periods and actual numeric benchmark cells inspected.", record_count=len(periods), numeric_value_count=numeric_cells, earliest_observation=min(periods) if periods else None, latest_observation=max(periods) if periods else None, observation_precision="month")
            evidence.check(name + ".freshness", "NOT_TESTED", "Historical benchmark workbook; latest period is reported without assuming it is current.", required=False)
            return
        # Unknown workbook layouts remain explicit; numeric cells alone are not
        # enough to claim that a market-price schema or geography was verified.
        evidence.check(name + ".records", "NOT_TESTED", "Workbook format verified; a reviewed sheet/header/measure profile is still required.")
        evidence.check(name + ".coverage", "NOT_TESTED", "Workbook transport and sheet names do not establish target geography or commodities.")
    finally:
        workbook.close()


def inspect_file(client, evidence, name, body, profile, *, complete=True):
    kind = profile.get("format", "").lower()
    if not complete and kind in {"zip", "xlsx"}:
        if not body.startswith(b"PK"):
            raise ProbeError("Expected ZIP/XLSX signature; the response is a different format.")
        evidence.check(name + ".signature", "PASS", "ZIP container prefix verified in a bounded sample; archive integrity is not claimed.")
        evidence.check(name + ".records", "NOT_TESTED", "Full archive was not read. Remove the source's sample_bytes limit within the existing byte/time budget, then validate the workbook schema.")
        evidence.check(name + ".completeness", "PARTIAL", "Only a file prefix was inspected.")
        return
    if kind == "pdf":
        if not body.startswith(b"%PDF-"):
            raise ProbeError("Expected PDF signature; a landing/error page was returned instead.")
        evidence.check(name + ".format", "PASS", "PDF signature verified; edition, table records and dates are not yet parsed.")
        evidence.check(name + ".records", "NOT_TESTED", "PDF table extraction profile has not been established.")
    elif kind == "xlsx":
        with safe_archive(body, client.max_bytes):
            pass
        inspect_workbook(evidence, name, body, profile)
    elif kind == "zip":
        with safe_archive(body, client.max_bytes) as archive:
            members = [n for n in archive.namelist() if n.lower().endswith(".xlsx") and not n.startswith("__MACOSX/")]
            evidence.check(name + ".archive", "PASS", "Archive opened in memory.", member_count=len(archive.namelist()), workbook_count=len(members))
            if not members:
                evidence.check(name + ".records", "NOT_TESTED", "No configured XLSX data resource found inside the archive.")
            for index, member in enumerate(members[:2]):
                inspect_workbook(evidence, name + f".workbook_{index + 1}", archive.read(member), profile)
            if len(members) > 2:
                evidence.check(name + ".other_workbooks", "NOT_TESTED", "Only the first two workbook members were inspected.", required=False)
    else:
        raise ProbeError("No reviewed file format is configured for this source.")


def probe_file(client, evidence, source, profile):
    body, _ = client.get("file_download", source["source_url"], read_limit=8192 if profile.get("format") == "pdf" else profile.get("sample_bytes"))
    inspect_file(client, evidence, "file", body, profile, complete=evidence.report["endpoints"][-1]["complete"])
    evidence.check("pagination", "NOT_APPLICABLE", "Configured source is a single downloadable file.")


def probe_nbs(client, evidence, source, profile):
    page, url = checked_html(client, evidence, "catalogue", source["source_url"])
    if "selectedfoodpricewatch" not in normalise(" ".join(page.text)):
        raise ProbeError("Catalogue identity does not match Selected Food Price Watch.")
    downloads = list(dict.fromkeys(urljoin(url, link["href"]) for link in page.links if re.search(r"/catalog/162/download/\d+", link["href"])))
    if not downloads:
        raise ProbeError("No actual data download links found in the catalogue response.")
    body, _ = client.get("first_listed_download", downloads[0], read_limit=profile.get("sample_bytes"))
    inspect_file(client, evidence, "first_listed_download", body, {"format": "zip", "workbook": "nbs_food_watch"}, complete=evidence.report["endpoints"][-1]["complete"])
    evidence.check("required_states", "NOT_TESTED", "National averages or state high/low extrema must not be mistaken for a complete four-state panel. Workbook state/price columns require a reviewed mapping.")
    evidence.check("other_downloads", "NOT_TESTED", "Only the first catalogue-listed download was tested; it is not assumed to be the newest observation period.", other_download_count=len(downloads) - 1, required=False)


def discover_naerls_requests(client, evidence, body, url):
    """Inspect bounded script resources; never execute JS or guess API paths."""
    from src.common.parsing import html_scripts

    blocks = html_scripts(body)
    scripts = [("inline", block["text"]) for block in blocks if not block["src"]]
    external = []
    for block in blocks:
        if not block["src"]:
            continue
        target = urljoin(url, block["src"])
        path = urlsplit(target).path.lower()
        if (urlsplit(target).scheme == "https" and urlsplit(target).netloc == urlsplit(url).netloc
                and "/libraries/" not in path
                and not any(library in path for library in ("jquery", "bootstrap", "datatable", "chart", "select2"))):
            external.append(target)
    external = list(dict.fromkeys(external))
    for index, target in enumerate(external[:2]):
        try:
            script, _ = client.get(f"market_script_{index + 1}", target)
            scripts.append((display_url(target), script.decode("utf-8-sig", "replace")))
        except ProbeError as error:
            evidence.check(f"market_script_{index + 1}.inspection", "NOT_TESTED", str(error))

    findings = []
    for origin, script in scripts:
        requests = []
        # Literal strings only. Concatenated paths are flagged as unresolved;
        # no request is issued to a discovered route in this diagnostic step.
        for match in re.finditer(r"\burl\s*:\s*(['\"])([^'\"\r\n]+)\1", script):
            literal = match.group(2)
            if literal.startswith(("/", "https://", "./", "../")):
                requests.append({"url_or_prefix": display_url(urljoin(url, literal)),
                                 "dynamic_suffix": bool(re.match(r"\s*\+", script[match.end():])),
                                 "method": "UNVERIFIED"})
        for match in re.finditer(r"\$\.(getJSON|get|post)\s*\(\s*(['\"])([^'\"\r\n]+)\2", script):
            requests.append({"url_or_prefix": display_url(urljoin(url, match.group(3))),
                             "dynamic_suffix": bool(re.match(r"\s*\+", script[match.end():])),
                             "method": "POST" if match.group(1) == "post" else "GET"})
        # Record key names only, never CSRF values or complete inline scripts.
        data_keys = sorted({key for match in re.finditer(r"\bdata\s*:\s*\{([^}]{0,2000})\}", script)
                            for key in re.findall(r"(?:^|,)\s*['\"]?([A-Za-z_][\w\[\]-]*)['\"]?\s*:", match.group(1))})
        selectors = sorted(set(re.findall(r"\$\(\s*['\"](#[A-Za-z_][\w-]*)['\"]\s*\)", script)))
        if requests or data_keys:
            findings.append({"script": origin, "request_literals": requests,
                             "payload_key_names": data_keys, "selectors": selectors})
    evidence.check("market_search.javascript_requests", "PARTIAL" if findings else "NOT_TESTED",
                   "Static script inspection only. Routes/prefixes are discovered, not verified data endpoints; methods, dependent values and responses need review before submission. Script code and tokens are not saved.",
                   findings=findings, external_scripts_inspected=min(2, len(external)),
                   external_scripts_uninspected=max(0, len(external) - 2))


def probe_naerls(client, evidence, source, profile):
    """Exercise one public form query, using only controls in the actual page."""
    body, url = client.get("market_search", source["source_url"])
    page = html_page(body)
    discover(evidence, "market_search", page, url)
    missing = [text for text in profile.get("expected_text", [])
               if normalise(text) not in normalise(" ".join(page.text))]
    evidence.check("market_search.identity", "FAIL" if missing else "PASS",
                   "Check the actual search page before submitting its form.", missing_identity_tokens=missing)
    if missing:
        return
    form = page.form("agricproduct")
    fields = {item.get("name") for item in form["inputs"] + form["selects"]}
    if not {"_token", "agricproduct", "priceregime", "market_id[]"}.issubset(fields):
        raise ProbeError("Expected NAERLS query controls are missing.")

    # Never record CSRF tokens, cookies or hidden input values in the evidence.
    public_fields = {"filter1", "filter2[]", "market_id[]", "agricproduct", "priceregime"}
    options = {}
    for select in form["selects"]:
        if select["name"] in public_fields:
            options[select["name"]] = [
                {"label": " ".join(option["text"]).strip(),
                 "value": option["value"] if option["value"] is not None else " ".join(option["text"])}
                for option in select["options"]
            ]
    evidence.check("market_search.options", "PASS", "Public options from the actual form; option presence is not price-record coverage.", options=options)
    evidence.check("required_location_coverage", "NOT_TESTED",
                   "Actual returned location fields must establish coverage; menu options alone cannot.",
                   locations={name: "UNVERIFIED" for name in ("Abuja/FCT", "Edo", "Benue", "Delta")})

    values = {}
    unavailable = []
    for field, choices in options.items():
        usable = [choice for choice in choices
                  if str(choice["value"]).strip() and not re.search(
                      r"^(?:[-\s]*select\b|[-\s]*choose\b)", choice["label"], re.I)]
        if not usable:
            unavailable.append(field)
        else:
            values[field] = usable[0]["value"]
    # Radio values must come from real controls, never invented enum constants.
    for field in public_fields - set(options):
        radios = [item for item in form["inputs"]
                  if item.get("name") == field and item.get("type", "").lower() == "radio"]
        if radios:
            selected = next((item for item in radios if "checked" in item), radios[0])
            if selected.get("value") is not None:
                values[field] = selected["value"]
    if not values.get("market_id[]") or not values.get("agricproduct") or not values.get("priceregime") or unavailable:
        try:
            values = fetch_naerls_selections(client, body, page, url)
        except ProbeError as error:
            discover_naerls_requests(client, evidence, body, url)
            evidence.check("market_query", "NOT_TESTED", str(error))
            return
    evidence.check("market_query.scope", "PASS",
                   "One bounded query using actual returned form options; no claim of all locations or commodities.", submitted_public_values=values)
    result, _ = post_form(client, "market_price_query", page, url, "_token", values)
    aliases = {
        "date": {"date", "pricedate", "observationdate", "period"},
        "commodity": {"commodity", "product", "agricproduct", "agriculturalproduct"},
        "price": {"price", "averageprice", "retailprice", "wholesaleprice"},
        "market": {"market", "marketname"},
        "state": {"state", "statename", "admin1", "adm1name"},
    }
    for table in result.tables:
        for index, row in enumerate(table):
            keys = [normalise(cell) for cell in row.cells]
            mapping = {name: next((key for key in keys if key in names), None)
                       for name, names in aliases.items()}
            if not all(mapping[name] for name in ("date", "commodity", "price")):
                continue
            if len(set(keys)) != len(keys) or not all(keys):
                raise ProbeError("Returned price-table headers are empty or repeated.")
            records = []
            for item in table[index + 1:]:
                if not any(cell.strip() for cell in item.cells) or item.cells == row.cells:
                    continue
                if len(item.cells) != len(keys):
                    raise ProbeError("Returned price table has inconsistent row widths.")
                records.append(dict(zip(keys, item.cells)))
            check_records(evidence, "market_query", records, date_field=mapping["date"],
                          price_fields=[mapping["price"]], commodity_field=mapping["commodity"],
                          location_field=mapping["market"] or mapping["state"],
                          max_age_days=profile.get("freshness_days"),
                          complete=client.evidence.report["endpoints"][-1]["complete"])
            if mapping["state"]:
                states = sorted({record[mapping["state"]] for record in records})
                matches = {label: any(normalise(state) in names for state in states)
                           for label, names in {
                               "Abuja/FCT": {"abuja", "fct", "abujafct", "federalcapitalterritory"},
                               "Edo": {"edo"}, "Benue": {"benue"}, "Delta": {"delta"}
                           }.items()}
                evidence.check("market_query.location_evidence", "PARTIAL",
                               "Locations present in this bounded query only; absent states remain unverified for the source.",
                               returned_states=states, locations_present_in_query=matches)
            evidence.check("market_query.coverage", "NOT_TESTED",
                           "One form result does not establish pagination, all markets, all crops or historical coverage.")
            return
    evidence.check("market_query.records", "NOT_TESTED",
                   "No supported date/commodity/price table identified. The response may contain no data, validation messages or a layout requiring additional parsing; HTTP success alone does not prove records.",
                   table_count=len(result.tables))


def probe_catalogue(client, evidence, source, profile):
    page, url = checked_html(client, evidence, "catalogue", source["source_url"])
    text = normalise(" ".join(page.text))
    expected = profile.get("expected_text", [])
    missing = [token for token in expected if normalise(token) not in text]
    evidence.check("catalogue_identity", "FAIL" if missing else "PASS", "Configured dataset identity checked against the actual response.", missing_identity_tokens=missing)
    if missing:
        raise ProbeError("The actual catalogue is a different dataset or a login/error page.")
    if profile.get("follow_microdata"):
        link = next((l for l in page.links if "get-microdata" in l["href"]), None)
        if link:
            access_page, _ = checked_html(client, evidence, "microdata_access", urljoin(url, link["href"]))
            evidence.check("microdata_access", "NOT_TESTED", "Access page inspected; a confirmed price-file request is still needed. No login or terms acceptance was submitted.")
    evidence.check("catalogue.records", "NOT_TESTED", "A catalogue/home page, form options or publication description is not verified price-record access.")
    evidence.check("observation_date", "NOT_TESTED", "Metadata/publication dates are not substituted for record observation dates.")
    evidence.check("pagination", "NOT_TESTED", "No verified data endpoint/pagination contract is configured.")


def probe_kamis(client, evidence, source, profile):
    for name, path in profile["endpoints"].items():
        try:
            page, _ = checked_html(client, evidence, name, urljoin(source["source_url"], path))
            headers, records = table_records(page, ["Market", "Commodity", "Wholesale", "Retail", "County", "Date"])
            for row in records:
                units = []
                for key in ("wholesale", "retail"):
                    value, separator, unit = row[key].partition("/")
                    row[key] = value.strip()
                    if separator:
                        units.append(unit.strip())
                row["price_unit"] = units[0] if units and len(set(units)) == 1 else ""
            check_records(evidence, name, records, date_field="date", price_fields=["wholesale", "retail"], commodity_field="commodity", location_field="county", unit_field="price_unit", max_age_days=profile.get("freshness_days"))
            evidence.check(name + ".target_crops", "PASS" if all(any(normalise(crop) in normalise(r.get("commodity")) for r in records) for crop in profile["commodities"]) else "NOT_TESTED", "Target crops checked only in the returned page; absence here does not establish absence from the complete source.")
        except ProbeError as exc:
            evidence.check(name + ".records", "FAIL", str(exc))
    evidence.check("pagination", "NOT_TESTED", "Server-side paging/date-search requests require an accessible, verified form/API contract.")


def probe_json_records(client, evidence, source, profile):
    """Exercise the shared fetcher and validate records in one JSON response."""
    records, _ = fetch_json_records(
        client, source["source_url"], params=profile.get("params"),
        records_path=profile.get("records_path", ""),
        record_shape=profile.get("record_shape", "list"),
    )
    if profile.get("validation_scope") == "structure":
        shapes = {}
        for row in records:
            shape = tuple(sorted(row))
            shapes[shape] = shapes.get(shape, 0) + 1
        evidence.check(
            "api.records", "PASS" if records else "PARTIAL",
            "JSON record objects retrieved; source fields are not renamed, converted or rejected for missing wholesale/retail fields. Empty results do not verify record availability.",
            record_count=len(records), validation_scope="structure",
            record_shapes=[{"fields": list(shape), "record_count": count} for shape, count in shapes.items()],
        )
        evidence.check(
            "price_semantics_and_freshness", "NOT_TESTED",
            "Price validity, basis, currency, units, date interpretation and freshness are not established by this structural access check.",
            required=False,
        )
        evidence.check(
            "pagination_and_coverage", "NOT_TESTED",
            "One summary response does not prove all markets, commodities, pages or historical coverage.",
            required=False,
        )
        return
    required = profile.get("required_fields", [])
    missing_counts = {key: sum(key not in row for row in records) for key in required}
    missing = sorted(key for key, count in missing_counts.items() if count)
    shapes = {}
    for row in records:
        shape = tuple(sorted(row))
        shapes[shape] = shapes.get(shape, 0) + 1
    examples = []
    for row in records:
        absent = [key for key in required if key not in row]
        if absent and len(examples) < 10:
            examples.append({
                "commodity": row.get(profile["commodity_field"]),
                "location": row.get(profile.get("location_field", "")),
                "missing_fields": absent,
                "available_fields": sorted(row),
            })
    evidence.check(
        "api.required_fields", "FAIL" if missing else "PASS",
        "Configured fields checked per record; diagnostics contain field names and commodity/location identifiers, not price values.",
        record_count=len(records), missing_fields=missing,
        missing_field_counts={key: missing_counts[key] for key in missing},
        record_shapes=[{"fields": list(shape), "record_count": count} for shape, count in shapes.items()],
        missing_field_examples=examples,
    )
    if missing:
        return
    complete = evidence.report["endpoints"][-1].get("complete", False)
    check_records(
        evidence, "api", records, date_field=profile["date_field"],
        price_fields=profile["price_fields"], commodity_field=profile["commodity_field"],
        location_field=profile.get("location_field"), unit_field=profile.get("unit_field"),
        max_age_days=profile.get("freshness_days"),
        required_locations=profile.get("required_locations"), complete=complete,
    )
    evidence.check("pagination", "NOT_TESTED", "This check covers one response only; no pagination or full-source coverage contract is verified.", required=False)
    evidence.check("source_scope", "NOT_TESTED", "A latest-summary response does not establish every market, commodity, historical date or reuse terms.", required=False)


def probe_json_pages(client, evidence, source, profile):
    records, complete = json_pages(client, evidence, "api", source["source_url"], records_path=profile["records_path"], next_path=profile["next_path"], max_pages=profile.get("max_pages", 3))
    check_records(evidence, "api", records, date_field=profile["date_field"], price_fields=profile["price_fields"], commodity_field=profile["commodity_field"], location_field=profile.get("location_field"), unit_field=profile.get("unit_field"), max_age_days=profile.get("freshness_days"), required_locations=profile.get("required_locations"), complete=complete)


def probe_cape_town(client, evidence, source, profile):
    page, _ = checked_html(client, evidence, "daily_table", source["source_url"])
    _, rows = table_records(page, ["DESC", "Container", "Mass", "Grade", "Low Price", "High Price", "Average Price"])
    match = re.search(r"Daily Statistical Prices\s*(?:for\s*)?[:\-]?\s*(\d{1,2}/\d{1,2}/\d{4})", " ".join(page.text), re.I)
    if not match:
        raise ProbeError("Daily-statistics observation-date label is missing; no retrieval date was used.")
    check_records(evidence, "daily_table", rows, date_field="date", date_override=match.group(1), price_fields=["lowprice", "highprice", "averageprice"], commodity_field="desc", max_age_days=profile.get("freshness_days"))
    price_order(evidence, "daily_table", rows, "lowprice", "highprice", "averageprice")
    evidence.check("pagination", "NOT_APPLICABLE", "Configured source returns one daily HTML price table.")


PROBES = {"naerls": probe_naerls, "durban": probe_durban, "tshwane": probe_tshwane, "joburg": probe_joburg, "strathmore": probe_strathmore, "csv": probe_csv, "file": probe_file, "nbs": probe_nbs, "catalogue": probe_catalogue, "kamis": probe_kamis, "json_records": probe_json_records, "json_pages": probe_json_pages, "cape_town": probe_cape_town}
