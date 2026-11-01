# Author: Daniel Wondyfiraw
# Updated: 2026-10-30
"""Assess an inventory with bounded concurrency; never collect price datasets.

Run: python -m src.batch_review --workers 4 --max-seconds 30
One worker handles each host sequentially. Do not run another tester concurrently.
"""

import argparse
import csv
import hashlib
import json
import threading
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from src.common.logging import get_logger
from src.common.source_config import GENERATED, load_candidates, load_testing_rules
from src.source_tester import assess, positive_seconds, save_result

logger = get_logger(__name__)
COLUMNS = ["source_id", "source_name", "url", "connection", "checks", "access_result",
           "format", "http_status", "latest_observation", "records_inspected", "issue"]


def review_row(result):
    """Summarise evidence without equating HTTP success with price access."""
    endpoints = result["evidence"].get("endpoints", [])
    checks = result["evidence"].get("checks", [])
    initial = next((e for e in endpoints if e["name"] == "initial_resource"), {})
    records = [c for c in checks if c["name"].endswith(".records")
               and c["status"] == "PASS" and c.get("record_count", 0) > 0]
    dates = sorted(str(c["latest_observation"]) for c in checks if c.get("latest_observation"))
    if initial.get("http_status") in {401, 403}:
        access = "ACCESS_DENIED_OR_AUTH_REQUIRED"
    elif initial.get("http_status") not in {None, 200}:
        access = "HTTP_REQUEST_FAILED"
    elif initial.get("tls") == "FAIL":
        access = "TLS_FAILED"
    elif result["connection_status"] != "PASS":
        error = (result.get("error") or initial.get("error") or "").lower()
        access = "DNS_FAILED" if "dns" in error else "TIMEOUT" if "timed out" in error or "time budget" in error else "CONNECTION_FAILED"
    elif records:
        access = "RECORDS_INSPECTED"
    elif any(c["name"].endswith((".format", ".archive")) and c["status"] == "PASS" for c in checks):
        access = "FILE_FORMAT_VERIFIED_SCHEMA_PENDING"
    elif result["profile"].get("response_format") == "html":
        access = "WEBPAGE_ONLY_DATA_UNVERIFIED"
    else:
        access = "RESOURCE_ACCESSIBLE_DATA_UNVERIFIED"
    issues = [result.get("error") or ""]
    issues.extend(e.get("error", "") for e in endpoints if e.get("status") != "PASS")
    issues.extend(c["name"] + ": " + c.get("detail", "") for c in checks if c["status"] == "FAIL")
    if not any(issues):
        issues.extend(c["detail"] for c in checks if c["status"] in {"PARTIAL", "NOT_TESTED"} and c.get("required", True))
    return dict(zip(COLUMNS, [
        result["source_id"], result["display_name"], result["source_url"],
        result["connection_status"], result["test_status"], access,
        result["profile"].get("response_format", "unknown"), initial.get("http_status"),
        dates[-1] if dates else "", max((c["record_count"] for c in records), default=0),
        " | ".join(dict.fromkeys(issue for issue in issues if issue)),
    ]))


def write_review(results):
    """Replace the human review CSV atomically; stable profile paths stay intact."""
    GENERATED.mkdir(parents=True, exist_ok=True)
    path = GENERATED / "source_review.csv"
    temporary = path.with_suffix(".csv.tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(review_row(r) for r in sorted(results, key=lambda r: r["source_id"]))
    temporary.replace(path)
    return path


def write_review_workbook(results):
    """Create a sharing copy; this workbook is not another input registry."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source review"
    fields = ["source_id", "source_name", "url", "access_result", "latest_observation", "records_inspected", "issue"]
    sheet.append(["Source ID", "Source", "URL", "Actual access result", "Latest observation", "Records inspected", "Remaining issue"])
    for result in sorted(results, key=lambda r: r["source_id"]):
        row = review_row(result)
        sheet.append([row[field] for field in fields])
        sheet.cell(sheet.max_row, 3).hyperlink = row["url"]
        sheet.cell(sheet.max_row, 3).style = "Hyperlink"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="14564B")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
        sheet.row_dimensions[row[0].row].height = 55
    for index, width in enumerate([40, 42, 65, 43, 22, 20, 85], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    path = GENERATED / "source_review.xlsx"
    workbook.save(path)
    return path


def worker_count(value):
    count = int(value)
    if not 1 <= count <= 8:
        raise argparse.ArgumentTypeError("Choose 1 to 8 workers.")
    return count


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="Default: latest generated urls_*.csv")
    parser.add_argument("--source-id", action="append", help="Optional source ID; repeat to select several")
    parser.add_argument("--workers", type=worker_count, default=4, help="Concurrent hosts, 1–8; default 4")
    parser.add_argument("--timeout", type=positive_seconds, default=6, help="Read timeout; connection timeout remains capped at 5 seconds")
    parser.add_argument("--max-seconds", type=positive_seconds, default=30, help="HTTP time budget per source; default 30 seconds")
    args = parser.parse_args(argv)
    sources, input_path = load_candidates(args.input)
    if args.source_id:
        missing = set(args.source_id) - {s["source_id"] for s in sources}
        if missing:
            parser.error("Unknown source IDs: " + ", ".join(sorted(missing)))
        sources = [s for s in sources if s["source_id"] in args.source_id]
    if not sources:
        parser.error("No sources selected.")
    urls = [s["source_url"] for s in sources]
    if len(set(urls)) != len(urls):
        parser.error("Duplicate URL strings in input; run preparation to consolidate them first.")
    rules = load_testing_rules()
    input_hash = hashlib.sha256(input_path.read_bytes()).hexdigest()
    groups = defaultdict(list)
    for source in sources:
        groups[urlsplit(source["source_url"]).hostname].append(source)
    results, lock = [], threading.Lock()
    started = datetime.now(timezone.utc).isoformat()

    def run_host(group):
        for source in group:
            pending = {**source, "connection_status": "NOT_TESTED", "test_status": "IN_PROGRESS",
                       "profile": dict(rules.get(source["source_id"], {})), "enabled_for_ingestion": False,
                       "attempt_started_at_utc": datetime.now(timezone.utc).isoformat()}
            save_result(pending, input_path, input_hash)
            result = assess(source, rules.get(source["source_id"], {}), args)
            result["attempt_started_at_utc"] = pending["attempt_started_at_utc"]
            save_result(result, input_path, input_hash)
            with lock:
                results.append(result)
                write_review(results)
                logger.info("%s/%s %s: connection=%s checks=%s", len(results), len(sources),
                            result["source_id"], result["connection_status"], result["test_status"])

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        for future in as_completed([pool.submit(run_host, group) for group in groups.values()]):
            future.result()
    links = [{"parent_source_id": r["source_id"], "url": url}
             for r in results for c in r["evidence"].get("checks", [])
             for url in c.get("discovered_data_links", []) if "?REDACTED" not in url]
    (GENERATED / "discovered_links.json").write_text(json.dumps(links, indent=2) + "\n", encoding="utf-8")
    counts = dict(Counter(review_row(r)["access_result"] for r in results))
    (GENERATED / "batch_review_report.json").write_text(json.dumps({
        "started_at_utc": started, "finished_at_utc": datetime.now(timezone.utc).isoformat(),
        "input_file": str(input_path), "input_sha256": input_hash,
        "sources_reviewed": len(results), "access_counts": counts,
        "scope": "Configured probes and initial-resource checks; discovered links are not automatically verified. No claim of complete endpoint or price coverage.",
        "workers": args.workers, "per_source_seconds": args.max_seconds,
    }, indent=2) + "\n", encoding="utf-8")
    logger.info("Review: %s", GENERATED / "source_review.csv")
    logger.info("Sharing workbook: %s", write_review_workbook(results))
    logger.info("Access results: %s", counts)
    return 0 if all(r["test_status"] == "PASS" for r in results) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, ValueError) as error:
        logger.error("Batch review could not finish: %s", error)
        raise SystemExit(1)
