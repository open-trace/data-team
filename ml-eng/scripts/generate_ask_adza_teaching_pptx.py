#!/usr/bin/env python3
"""Generate the Ask ADZA community teaching PowerPoint (16:9)."""
from __future__ import annotations

from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

from ask_adza_teaching_content import DOC_VERSION, DOCS_DIR, OUTPUT_PPTX, SLIDES

# Widescreen 16:9 — same chrome as generate_rag_architecture_pptx.py
_SLIDE_W = Inches(13.333)
_SLIDE_H = Inches(7.5)

_BG = RGBColor(0x0F, 0x1A, 0x14)
_BG_ALT = RGBColor(0x16, 0x24, 0x1C)
_ACCENT = RGBColor(0x2F, 0x9E, 0x6B)
_ACCENT_WARM = RGBColor(0xC4, 0xA3, 0x5A)
_TEXT = RGBColor(0xF2, 0xF0, 0xE9)
_MUTED = RGBColor(0xA8, 0xB5, 0xAC)
_CARD = RGBColor(0x1A, 0x2B, 0x22)
_CARD_LINE = RGBColor(0x2A, 0x3F, 0x33)
_HEADER_CELL = RGBColor(0x14, 0x28, 0x1E)


def _set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, _SLIDE_W, _SLIDE_H)
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    fill.line.fill.background()
    spTree = slide.shapes._spTree
    sp = fill._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _accent_bar(slide, top=Inches(0)) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, _SLIDE_W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _ACCENT
    bar.line.fill.background()


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 18,
    bold: bool = False,
    color: RGBColor = _TEXT,
    align=PP_ALIGN.LEFT,
    font_name: str = "Calibri",
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def _add_bullets(
    slide,
    left,
    top,
    width,
    height,
    bullets: list[str],
    *,
    size: int = 16,
    color: RGBColor = _TEXT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets[:8]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {bullet}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"


def _set_notes(slide, text: str | None) -> None:
    notes = (text or "").strip()
    if not notes:
        return
    slide.notes_slide.notes_text_frame.text = notes


def _title_bar(slide, title: str) -> None:
    _accent_bar(slide)
    _add_text_box(
        slide,
        Inches(0.7),
        Inches(0.28),
        Inches(12),
        Inches(0.65),
        title,
        size=24,
        bold=True,
        color=_TEXT,
    )


def _footnote(slide, text: str | None) -> None:
    if not text:
        return
    _add_text_box(
        slide,
        Inches(0.7),
        Inches(6.85),
        Inches(12),
        Inches(0.4),
        text,
        size=12,
        color=_MUTED,
    )


def _rounded_card(slide, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = _CARD
    card.line.color.rgb = _CARD_LINE
    return card


def _set_cell_fill(cell, color: RGBColor) -> None:
    hex_val = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        tag = child.tag.lower()
        if tag.endswith("solidfill") or tag.endswith("gradfill") or tag.endswith("nofill"):
            tcPr.remove(child)
    solid = parse_xml(
        f'<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:srgbClr val="{hex_val}"/></a:solidFill>'
    )
    tcPr.append(solid)


def _set_cell_text(
    cell,
    text: str,
    *,
    size: int,
    bold: bool = False,
    color: RGBColor = _TEXT,
) -> None:
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _render_title(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG)
    _accent_bar(slide, Inches(2.15))
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(1.55),
        Inches(11.5),
        Inches(0.4),
        slide_spec.get("kicker", ""),
        size=14,
        color=_ACCENT,
    )
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(2.35),
        Inches(11.5),
        Inches(1.1),
        slide_spec["title"],
        size=40,
        bold=True,
        color=_TEXT,
    )
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(3.55),
        Inches(11.5),
        Inches(0.5),
        slide_spec.get("subtitle", ""),
        size=18,
        color=_MUTED,
    )
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(4.25),
        Inches(11.5),
        Inches(0.5),
        slide_spec.get("footer", ""),
        size=14,
        color=_ACCENT_WARM,
    )
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(6.6),
        Inches(11.5),
        Inches(0.4),
        f"Generated {date.today().isoformat()}  ·  content v{DOC_VERSION}  ·  ml/rag",
        size=12,
        color=_MUTED,
    )
    _set_notes(slide, slide_spec.get("notes"))


def _render_section(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG_ALT)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.9), Inches(1.2), Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = _ACCENT
    bar.line.fill.background()
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(3.2),
        Inches(11.5),
        Inches(1),
        slide_spec["title"],
        size=32,
        bold=True,
        color=_TEXT,
    )
    _add_text_box(
        slide,
        Inches(0.9),
        Inches(4.25),
        Inches(11.5),
        Inches(0.7),
        slide_spec.get("subtitle", ""),
        size=16,
        color=_MUTED,
    )
    _set_notes(slide, slide_spec.get("notes"))


def _render_bullets(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG)
    _title_bar(slide, slide_spec["title"])
    _add_bullets(
        slide,
        Inches(0.8),
        Inches(1.15),
        Inches(11.7),
        Inches(5.4),
        list(slide_spec.get("bullets") or []),
        size=17,
    )
    _footnote(slide, slide_spec.get("footnote"))
    _set_notes(slide, slide_spec.get("notes"))


def _render_two_column(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG)
    _title_bar(slide, slide_spec["title"])
    for left, card_title, bullets in (
        (Inches(0.55), slide_spec["left_title"], list(slide_spec.get("left") or [])),
        (Inches(6.85), slide_spec["right_title"], list(slide_spec.get("right") or [])),
    ):
        _rounded_card(slide, left, Inches(1.15), Inches(5.9), Inches(5.45))
        _add_text_box(
            slide,
            left + Inches(0.25),
            Inches(1.35),
            Inches(5.4),
            Inches(0.45),
            card_title,
            size=16,
            bold=True,
            color=_ACCENT,
        )
        _add_bullets(
            slide,
            left + Inches(0.25),
            Inches(1.9),
            Inches(5.4),
            Inches(4.4),
            bullets,
            size=14,
        )
    _footnote(slide, slide_spec.get("footnote"))
    _set_notes(slide, slide_spec.get("notes"))


def _render_cards(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG)
    _title_bar(slide, slide_spec["title"])
    cards = list(slide_spec.get("cards") or [])[:4]
    n = max(1, len(cards))
    gap = Inches(0.25)
    margin = Inches(0.55)
    total_w = _SLIDE_W - margin * 2
    card_w = int((total_w - gap * (n - 1)) / n)
    top = Inches(1.2)
    height = Inches(5.35)
    for i, card in enumerate(cards):
        left = int(margin + i * (card_w + gap))
        _rounded_card(slide, left, top, card_w, height)
        _add_text_box(
            slide,
            left + Inches(0.2),
            Inches(1.45),
            card_w - Inches(0.4),
            Inches(0.7),
            str(card.get("title") or ""),
            size=16,
            bold=True,
            color=_ACCENT,
        )
        _add_text_box(
            slide,
            left + Inches(0.2),
            Inches(2.25),
            card_w - Inches(0.4),
            Inches(3.9),
            str(card.get("body") or ""),
            size=14,
            color=_TEXT,
        )
    _footnote(slide, slide_spec.get("footnote"))
    _set_notes(slide, slide_spec.get("notes"))


def _render_table(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG)
    _title_bar(slide, slide_spec["title"])
    headers = list(slide_spec.get("headers") or [])
    rows = list(slide_spec.get("rows") or [])
    if not headers:
        _set_notes(slide, slide_spec.get("notes"))
        return
    n_cols = len(headers)
    n_rows = 1 + len(rows)
    left = Inches(0.5)
    top = Inches(1.15)
    width = Inches(12.3)
    height = Inches(5.45)
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table
    col_w = int(width / n_cols)
    for i in range(n_cols):
        table.columns[i].width = col_w

    header_size = 12 if n_cols >= 4 else 13
    body_size = 11 if n_cols >= 4 or any(len(r) > 3 for r in rows) else 12
    long_cell = max((len(str(c)) for row in rows for c in row), default=0)
    if long_cell > 70:
        body_size = 11
        header_size = 12

    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        _set_cell_fill(cell, _HEADER_CELL)
        _set_cell_text(cell, str(header), size=header_size, bold=True, color=_ACCENT)
    for r, row in enumerate(rows, start=1):
        for c in range(n_cols):
            cell = table.cell(r, c)
            _set_cell_fill(cell, _CARD)
            value = row[c] if c < len(row) else ""
            _set_cell_text(cell, str(value), size=body_size, color=_TEXT)

    _footnote(slide, slide_spec.get("footnote"))
    _set_notes(slide, slide_spec.get("notes"))


def _render_pipeline(prs: Presentation, slide_spec: dict) -> None:
    slide = _blank(prs)
    _set_slide_bg(slide, _BG)
    _title_bar(slide, slide_spec["title"])
    steps = [str(s) for s in (slide_spec.get("steps") or [])][:8]
    n = max(1, len(steps))
    margin = Inches(0.45)
    gap = Inches(0.12)
    total_w = _SLIDE_W - margin * 2
    box_w = int((total_w - gap * (n - 1)) / n)
    top = Inches(1.35)
    height = Inches(1.55)
    for i, step in enumerate(steps):
        left = int(margin + i * (box_w + gap))
        _rounded_card(slide, left, top, box_w, height)
        _add_text_box(
            slide,
            left + Inches(0.08),
            Inches(1.42),
            box_w - Inches(0.16),
            Inches(0.32),
            f"{i + 1:02d}",
            size=11,
            bold=True,
            color=_ACCENT,
            align=PP_ALIGN.CENTER,
        )
        _add_text_box(
            slide,
            left + Inches(0.08),
            Inches(1.78),
            box_w - Inches(0.16),
            Inches(0.95),
            step,
            size=13 if n <= 6 else 11,
            bold=True,
            color=_TEXT,
            align=PP_ALIGN.CENTER,
        )
        if i < n - 1:
            chev_left = left + box_w - Inches(0.02)
            _add_text_box(
                slide,
                chev_left,
                Inches(1.85),
                gap + Inches(0.08),
                Inches(0.4),
                "→",
                size=14,
                color=_ACCENT,
                align=PP_ALIGN.CENTER,
            )
    caption = str(slide_spec.get("caption") or "")
    if caption:
        _rounded_card(slide, Inches(0.55), Inches(3.2), Inches(12.2), Inches(3.35))
        _add_text_box(
            slide,
            Inches(0.85),
            Inches(3.45),
            Inches(11.6),
            Inches(2.9),
            caption,
            size=16,
            color=_TEXT,
        )
    _footnote(slide, slide_spec.get("footnote"))
    _set_notes(slide, slide_spec.get("notes"))


_RENDERERS = {
    "title": _render_title,
    "section": _render_section,
    "bullets": _render_bullets,
    "two_column": _render_two_column,
    "cards": _render_cards,
    "table": _render_table,
    "pipeline": _render_pipeline,
}


def build_pptx(output=OUTPUT_PPTX):
    prs = Presentation()
    prs.slide_width = _SLIDE_W
    prs.slide_height = _SLIDE_H
    for spec in SLIDES:
        kind = str(spec.get("kind") or "")
        renderer = _RENDERERS.get(kind)
        if renderer is None:
            raise SystemExit(f"Unknown slide kind: {kind!r}")
        renderer(prs, spec)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


def main() -> None:
    path = build_pptx(OUTPUT_PPTX)
    print(f"Wrote {path}  ({len(SLIDES)} slides, v{DOC_VERSION})")


if __name__ == "__main__":
    main()
